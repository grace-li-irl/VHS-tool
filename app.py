import streamlit as st
import pandas as pd
import openpyxl
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border
from openpyxl.utils import get_column_letter
from io import BytesIO
import re
import copy
import os

st.set_page_config(page_title="多企业表格处理工具", layout="centered")
st.title("🏢 多企业表格处理工具")

TEMPLATE_FILE = "都会养老分层20260424.xlsx"
if not os.path.exists(TEMPLATE_FILE):
    st.error(f"❌ 样式模板文件 {TEMPLATE_FILE} 未找到！请将模板文件放置在程序同一目录下。")
    st.stop()

# ==================== 套餐等级映射表 ====================
PACKAGE_MAPPING = {
    "新客户养老分层-经典套餐": "大都会养老分层经典套餐",
    "新客户养老分层-白金套餐": "大都会养老分层白金套餐",
    "新客户养老分层-贵宾套餐": "大都会养老分层贵宾套餐",
    "都会金彩养老年金保险": "大都会都会金彩套餐",
    "都汇康健": "都会安泰服务套餐",
    "都会安泰": "都会安泰服务套餐",
    "都会传家（A款）终身寿险（分红型）": "大都会都会传家悦享版标准套餐",
    "都会颐年（2023）养老年金保险（分红型）": "大都会都会颐年2023养老年金套餐",
    "都会颐年（2024）养老年金保险（分红型）": "大都会都会颐年2024养老年金套餐",
    "都会颐年（2024）养老年金保险（分红型）-个人养老金": "大都会都会颐年2024个人养老金套餐",
    "都会颐年养老年金保险（分红型）（AGY）": "都会颐年养老年金套餐",
    "都会颐享两全保险（分红型）（常规版）": "大都会都会颐享常规版套餐",
    "都会颐享两全保险（分红型）（个人养老金版）": "大都会都会颐享个人养老金版套餐",
    "都会长虹(BXS)": "大都会都会长虹套餐",
    "都会长青(BXS)": "大都会都会长青套餐",
    "都会长青(DMTM)": "大都会都会长青套餐",
    "都会长盈年金保险（分红型）": "大都会都会长盈套餐",
    "倍佑一生（优享版）重大疾病保险改成": "大都会倍佑一生套餐",
    "倍佑一生重大疾病保险(DMTM)改成": "大都会倍佑一生套餐"
}

# ==================== 辅助函数 ====================
def add_prefix(prefix):
    return lambda x: f"{prefix}{x}" if pd.notna(x) else x

def remove_star(x):
    return str(x).replace('*', '') if pd.notna(x) else x

def last_six(x):
    s = str(x)
    return s[-6:] if len(s) >= 6 else s

def gender_fm(x):
    if pd.isna(x):
        return ''
    s = str(x).upper()
    if s == 'F':
        return '女'
    elif s == 'M':
        return '男'
    return s

def to_text(x):
    if pd.isna(x):
        return ''
    return str(x)

def pass_through(x):
    return x

def remove_spaces(x):
    if pd.isna(x):
        return ''
    return str(x).replace(' ', '')

fixed_output_columns = [
    "证件号码", "部门", "客户ID", "员工姓名", "手机号码", "性别", "生日", "密码",
    "是否需要绑定手机(0需要；1不需要)", "套餐名称", "保单生效时间",
    "服务开始时间", "服务结束时间", "是否是企业管理者"
]

# ==================== 企业规则（支持列名别名） ====================
COMPANY_RULES = {
    "大都会": {
        "source_columns": ["证件编号", "分公司", "客户编码", "客户姓名", "性别", "出生日期", "套餐等级", "加入养老产品增值服务日期"],
        "field_map": {
            "证件号码": (["证件编号", "证件号码", "证件号"], lambda x: to_text(remove_star(x))),
            "部门": (["分公司", "部门"], pass_through),
            "客户ID": (["客户编码", "客户号","客户ID"], lambda x: x if str(x).upper().startswith("MET") else f"MET{x}"),
            "员工姓名": (["客户姓名", "员工姓名", "姓名"], remove_spaces),
            "性别": (["性别", "客户性别"], pass_through),
            "生日": (["出生日期", "生日"], pass_through),
            "密码": (None, "666666"),
            "是否需要绑定手机(0需要；1不需要)": (None, 0),
            "套餐名称": (["套餐等级", "产品/服务名称", "套餐名称", "主险产品名称"], lambda x: PACKAGE_MAPPING.get(x, x)),
            "保单生效时间": (["加入养老产品增值服务日期", "保单生效时间", "服务生效日", "保单生效日"], pass_through),
            "服务开始时间": (["加入养老产品增值服务日期", "服务开始时间", "保单生效时间", "保单生效日"], pass_through),
            "服务结束时间": (None, "2028-05-17"),
            "是否是企业管理者": (None, "否")
        }
    },
    "安盛": {
        "source_columns": ["证件号", "产品名称", "客户号 （Party ID）", "姓名中文", "性别", "服务开始时间", "服务到期时间"],
        "field_map": {
            "证件号码": (["证件号"], to_text),
            "部门": (["产品名称"], pass_through),
            "客户ID": (["客户号 （Party ID）"], add_prefix("AXA")),
            "员工姓名": (["姓名中文"], remove_spaces),
            "性别": (["性别"], pass_through),
            "生日": (None, ""),
            "密码": (None, "666666"),
            "是否需要绑定手机(0需要；1不需要)": (None, 0),
            "套餐名称": (None, "2026安盛心理咨询套餐"),
            "保单生效时间": (["服务开始时间"], pass_through),
            "服务开始时间": (["服务开始时间"], pass_through),
            "服务结束时间": (["服务到期时间"], pass_through),
            "是否是企业管理者": (None, "否")
        }
    },
    "安永": {
        "source_columns": ["工号", "城市", "中文名", "性别"],
        "field_map": {
            "证件号码": (["工号"], to_text),
            "部门": (["城市"], pass_through),
            "客户ID": (["工号"], add_prefix("EY")),
            "员工姓名": (["中文名"], remove_spaces),
            "性别": (["性别"], gender_fm),
            "生日": (None, ""),
            "密码": (None, "666666"),
            "是否需要绑定手机(0需要；1不需要)": (None, 0),
            "套餐名称": (None, "EY安永EAP套餐"),
            "保单生效时间": (None, "2024-10-01"),
            "服务开始时间": (None, "2024-10-01"),
            "服务结束时间": (None, "2027-09-30"),
            "是否是企业管理者": (None, "否")
        }
    },
    "富国基金": {
        "source_columns": ["联系手机", "部门", "姓名", "性别"],
        "field_map": {
            "证件号码": (["联系手机"], to_text),
            "部门": (["部门"], pass_through),
            "客户ID": (["联系手机"], add_prefix("fg")),
            "员工姓名": (["姓名"], remove_spaces),
            "性别": (["性别"], gender_fm),
            "生日": (None, ""),
            "密码": (["联系手机"], last_six),
            "是否需要绑定手机(0需要；1不需要)": (None, 1),
            "套餐名称": (None, "富国基金HEAP套餐"),
            "保单生效时间": (None, "2026-01-01"),
            "服务开始时间": (None, "2026-01-01"),
            "服务结束时间": (None, "2026-12-31"),
            "是否是企业管理者": (None, "否")
        }
    },
    "晖致": {
        "source_columns": ["手机号", "服务部", "联系手机"],
        "field_map": {
            "证件号码": (["手机号"], to_text),
            "部门": (["服务部"], pass_through),
            "客户ID": (["手机号"], add_prefix("viatris-")),
            "员工姓名": (["手机号"], remove_spaces),
            "性别": (None, ""),
            "生日": (None, ""),
            "密码": (["联系手机"], last_six),
            "是否需要绑定手机(0需要；1不需要)": (None, 1),
            "套餐名称": (None, "晖致医药HEAP套餐"),
            "保单生效时间": (None, "2025-10-01"),
            "服务开始时间": (None, "2025-10-01"),
            "服务结束时间": (None, "2026-10-01"),
            "是否是企业管理者": (None, "否")
        }
    },
    "金士顿": {
        "source_columns": ["工号", "部门"],
        "field_map": {
            "证件号码": (["工号"], to_text),
            "部门": (["部门"], pass_through),
            "客户ID": (["工号"], add_prefix("SHKING")),
            "员工姓名": (["工号"], lambda x: add_prefix("SHKING")(remove_spaces(x))),
            "性别": (None, ""),
            "生日": (None, ""),
            "密码": (None, "666666"),
            "是否需要绑定手机(0需要；1不需要)": (None, 1),
            "套餐名称": (None, "上海金士顿HEAP套餐"),
            "保单生效时间": (None, "2026-01-01"),
            "服务开始时间": (None, "2026-01-01"),
            "服务结束时间": (None, "2026-10-31"),
            "是否是企业管理者": (None, "否")
        }
    },
    "聚物腾云": {
        "source_columns": ["手机号", "服务部"],
        "field_map": {
            "证件号码": (["手机号"], to_text),
            "部门": (["服务部"], pass_through),
            "客户ID": (["手机号"], add_prefix("Altium-")),
            "员工姓名": (None, ""),
            "性别": (None, ""),
            "生日": (None, ""),
            "密码": (["手机号"], pass_through),
            "是否需要绑定手机(0需要；1不需要)": (None, 1),
            "套餐名称": (None, "聚物腾云HEAP套餐"),
            "保单生效时间": (None, "2026-06-05"),
            "服务开始时间": (None, "2026-06-05"),
            "服务结束时间": (None, "2027-06-04"),
            "是否是企业管理者": (None, "否")
        }
    },
    "深圳医科院": {
        "source_columns": ["学号", "服务部", "姓名"],
        "field_map": {
            "证件号码": (["学号"], to_text),
            "部门": (["服务部"], pass_through),
            "客户ID": (["学号"], add_prefix("SMART")),
            "员工姓名": (["姓名"], remove_spaces),
            "性别": (None, ""),
            "生日": (None, ""),
            "密码": (["学号"], pass_through),
            "是否需要绑定手机(0需要；1不需要)": (None, 0),
            "套餐名称": (None, "深圳医科院EAP套餐"),
            "保单生效时间": (None, "2026-01-01"),
            "服务开始时间": (None, "2026-01-01"),
            "服务结束时间": (None, "2027-01-01"),
            "是否是企业管理者": (None, "否")
        }
    },
    "微盟电子": {
        "source_columns": ["行動電話", "服務部", "員工編號", "員工姓名", "性別"],
        "field_map": {
            "证件号码": (["行動電話"], to_text),
            "部门": (["服務部"], pass_through),
            "客户ID": (["員工編號"], add_prefix("MSIK")),
            "员工姓名": (["員工姓名"], remove_spaces),
            "性别": (["性別"], gender_fm),
            "生日": (None, ""),
            "密码": (None, "666666"),
            "是否需要绑定手机(0需要；1不需要)": (None, 1),
            "套餐名称": (None, "微盟电子HEAP套餐"),
            "保单生效时间": (None, "2025-09-26"),
            "服务开始时间": (None, "2025-09-26"),
            "服务结束时间": (None, "2026-09-25"),
            "是否是企业管理者": (None, "否")
        }
    },
    "盐田国际": {
        "source_columns": ["工号", "姓名"],
        "field_map": {
            "证件号码": (["工号"], to_text),
            "部门": (None, "服务部"),
            "客户ID": (["工号"], add_prefix("YICT")),
            "员工姓名": (["姓名"], remove_spaces),
            "性别": (None, ""),
            "生日": (None, ""),
            "密码": (None, "666666"),
            "是否需要绑定手机(0需要；1不需要)": (None, 0),
            "套餐名称": (None, "盐田国际EAP项目"),
            "保单生效时间": (None, "2025-01-01"),
            "服务开始时间": (None, "2025-01-01"),
            "服务结束时间": (None, "2026-12-31"),
            "是否是企业管理者": (None, "否")
        }
    }
}

company_list = list(COMPANY_RULES.keys())
selected_company = st.selectbox("🏭 选择企业", company_list)

# ==================== 处理函数 ====================
def process_data(df_source, rules):
    df_result = pd.DataFrame(index=df_source.index)
    for target_col, (sources, func) in rules["field_map"].items():
        if sources is None:
            df_result[target_col] = func
            continue
        if isinstance(sources, str):
            sources = [sources]
        found_col = None
        for col in sources:
            if col in df_source.columns:
                found_col = col
                break
        if found_col is None:
            df_result[target_col] = ''
            continue
        series = df_source[found_col]
        if func and callable(func):
            series = series.apply(func)
        df_result[target_col] = series
    for col in fixed_output_columns:
        if col not in df_result.columns:
            df_result[col] = ''
    return df_result

def extract_date_from_filename(filename):
    patterns = [
        r'(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})(?:日)?',
        r'(\d{4})(\d{2})(\d{2})'
    ]
    for pat in patterns:
        match = re.search(pat, filename)
        if match:
            groups = match.groups()
            if len(groups) == 3:
                year, month, day = groups
                return f"{year}{int(month):02d}{int(day):02d}"
    from datetime import datetime
    return datetime.now().strftime("%Y%m%d")

def set_cell_style(cell, style_dict):
    if style_dict.get("font"):
        cell.font = style_dict["font"]
    if style_dict.get("fill"):
        cell.fill = style_dict["fill"]
    if style_dict.get("alignment"):
        cell.alignment = style_dict["alignment"]
    if style_dict.get("border"):
        cell.border = style_dict["border"]
    if style_dict.get("number_format"):
        cell.number_format = style_dict["number_format"]

# ==================== 文件上传（添加 key 以实现企业切换时重置） ====================
uploaded_file = st.file_uploader("📂 上传原始数据文件（Excel 或 CSV）", 
                                 type=["xlsx", "xls", "csv"],
                                 key=selected_company)   # <--- 关键修改

if uploaded_file and selected_company:
    if st.button("🚀 开始处理", type="primary"):
        with st.spinner("正在处理数据，请稍候..."):
            try:
                if uploaded_file.name.endswith('.csv'):
                    try:
                        df_raw = pd.read_csv(uploaded_file, encoding='utf-8-sig')
                    except UnicodeDecodeError:
                        df_raw = pd.read_csv(uploaded_file, encoding='gbk')
                else:
                    df_raw = pd.read_excel(uploaded_file)
                st.success(f"✅ 原始数据读取成功，共 {len(df_raw)} 行")
                with st.expander("查看原始数据前5行"):
                    st.dataframe(df_raw.head())

                rules = COMPANY_RULES[selected_company]
                missing_cols = [col for col in rules["source_columns"] if col not in df_raw.columns]
                if missing_cols:
                    st.warning(f"原始数据缺少以下建议列：{missing_cols}（可能通过其他别名找到）")

                df_proc = process_data(df_raw, rules)
                df_proc = df_proc[fixed_output_columns]

                # 加载模板样式
                wb_template = load_workbook(TEMPLATE_FILE)
                ws_template = wb_template.active
                header_row = [cell.value for cell in ws_template[1] if cell.value is not None]
                template_order = [col for col in header_row if col in df_proc.columns]
                for col in fixed_output_columns:
                    if col not in template_order:
                        template_order.append(col)
                df_proc = df_proc[template_order]

                col_widths = {}
                for col_letter, dim in ws_template.column_dimensions.items():
                    if dim.width:
                        col_widths[col_letter] = dim.width
                header_styles = {}
                for col_idx in range(1, ws_template.max_column + 1):
                    src_cell = ws_template.cell(row=1, column=col_idx)
                    header_styles[col_idx] = {
                        "font": copy.copy(src_cell.font) if src_cell.font else Font(),
                        "fill": copy.copy(src_cell.fill) if src_cell.fill else PatternFill(),
                        "alignment": copy.copy(src_cell.alignment) if src_cell.alignment else Alignment(),
                        "border": copy.copy(src_cell.border) if src_cell.border else Border(),
                        "number_format": src_cell.number_format
                    }
                row_height = ws_template.row_dimensions[1].height if ws_template.row_dimensions[1].height else None

                wb_new = Workbook()
                ws_new = wb_new.active
                ws_new.title = "处理结果"

                for col_idx, col_name in enumerate(template_order, start=1):
                    ws_new.cell(row=1, column=col_idx, value=col_name)

                for row_idx, row_data in enumerate(df_proc.values.tolist(), start=2):
                    for col_idx, value in enumerate(row_data, start=1):
                        cell = ws_new.cell(row=row_idx, column=col_idx, value=value)
                        if col_idx == 1:
                            cell.number_format = '@'

                new_col_cnt = len(template_order)
                orig_letters = list(col_widths.keys())
                for i in range(new_col_cnt):
                    col_letter = get_column_letter(i + 1)
                    if i < len(orig_letters):
                        ws_new.column_dimensions[col_letter].width = col_widths[orig_letters[i]]
                    else:
                        ws_new.column_dimensions[col_letter].width = 10

                # 调整服务结束时间列宽与保单生效时间一致
                try:
                    policy_start_col = "保单生效时间"
                    service_end_col = "服务结束时间"
                    if policy_start_col in template_order and service_end_col in template_order:
                        policy_idx = template_order.index(policy_start_col) + 1
                        end_idx = template_order.index(service_end_col) + 1
                        policy_width = ws_new.column_dimensions[get_column_letter(policy_idx)].width
                        if policy_width:
                            ws_new.column_dimensions[get_column_letter(end_idx)].width = policy_width
                except Exception:
                    pass

                for col_idx in range(1, new_col_cnt + 1):
                    cell = ws_new.cell(row=1, column=col_idx)
                    style_idx = col_idx if col_idx in header_styles else 1
                    style = header_styles[style_idx]
                    set_cell_style(cell, style)
                if row_height:
                    ws_new.row_dimensions[1].height = row_height

                data_font = Font(name='等线', size=11)
                for row in range(2, len(df_proc) + 2):
                    for col in range(1, new_col_cnt + 1):
                        cell = ws_new.cell(row=row, column=col)
                        cell.font = data_font

                target_col_name = "是否需要绑定手机(0需要；1不需要)"
                if target_col_name in template_order:
                    col_idx = template_order.index(target_col_name) + 1
                    for row in range(2, len(df_proc) + 2):
                        cell = ws_new.cell(row=row, column=col_idx)
                        cell.alignment = Alignment(horizontal='left')

                output = BytesIO()
                wb_new.save(output)
                output.seek(0)

                date_str = extract_date_from_filename(uploaded_file.name)
                output_filename = f"{selected_company}{date_str}.xlsx"
                st.success(f"✅ 处理完成！共输出 {len(df_proc)} 行数据。")
                st.download_button(
                    label="📥 下载处理后的 Excel",
                    data=output,
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            except Exception as e:
                st.error(f"❌ 处理出错：{e}")
                import traceback
                st.code(traceback.format_exc())
